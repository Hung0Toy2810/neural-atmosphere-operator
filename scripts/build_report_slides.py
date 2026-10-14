"""Build the Stage 1/Stage 2 project presentation from the supplied course template."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.presentation import Presentation as PresentationType
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT.parent / "LuongNguyenThanhHung_final.pptx"
OUTPUT = ROOT / "reports" / "NeuralAtmosphereOperator_Stage1_Stage2.pptx"

NAVY = "20384F"
BLUE = "1976D2"
BLUE_2 = "2F95DE"
PALE = "EEF5FB"
PALE_2 = "F7FAFD"
GREEN = "2E8B57"
GREEN_PALE = "EAF6EF"
RED = "D94C4C"
RED_PALE = "FFF0F0"
ORANGE = "E28A2B"
GRAY = "5B6773"
MID = "A9B7C6"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def clear_slides(prs: PresentationType) -> None:
    for slide_id in list(prs.slides._sldIdLst):  # noqa: SLF001 - python-pptx API gap
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)  # noqa: SLF001


def set_fill(shape, color: str, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency


def set_line(shape, color: str, width: float = 1.0) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def text_box(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = NAVY,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.04,
    font: str = "Arial",
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def rich_text(slide, lines: list[tuple[str, float, str, bool]], x, y, w, h):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(0.05)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    for index, (value, size, color, bold) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(7)
        run = paragraph.add_run()
        run.text = value
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def rounded_box(slide, x, y, w, h, *, fill=PALE, line=MID, radius=True):
    kind = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = 0.12
    set_fill(shape, fill)
    set_line(shape, line, 1.0)
    return shape


def add_title(slide, title: str, number: int, subtitle: str | None = None) -> None:
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.35), Inches(0.17), Inches(12.62), Inches(0.055)
    )
    set_fill(line, BLUE)
    line.line.fill.background()
    text_box(slide, title, 0.42, 0.28, 10.7, 0.48, size=25, bold=True)
    text_box(slide, f"{number:02d}", 12.25, 0.31, 0.45, 0.25, size=10, color=BLUE, bold=True, align=PP_ALIGN.RIGHT)
    if subtitle:
        text_box(slide, subtitle, 0.45, 0.73, 11.7, 0.3, size=10.5, color=GRAY)


def add_footer(slide) -> None:
    text_box(slide, "Lương Nguyễn Thành Hưng · NeuralAtmosphereOperator", 0.42, 7.18, 5.4, 0.18, size=6.5, color="8A96A3")
    text_box(slide, "Stage 1 → Stage 2", 11.35, 7.18, 1.45, 0.18, size=6.5, color="8A96A3", align=PP_ALIGN.RIGHT)


def add_bullets(
    slide,
    items: Iterable[str],
    x,
    y,
    w,
    h,
    *,
    size: float = 16,
    color=NAVY,
    spacing=8,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.text = "•  " + paragraph.text
        paragraph.space_after = Pt(spacing)
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = rgb(color)
    return shape


def connector(slide, x1, y1, x2, y2, color=BLUE, width=2.0):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    set_line(shape, color, width)
    line = shape._element.spPr.get_or_add_ln()
    arrow = OxmlElement('a:tailEnd')
    arrow.set('type', 'triangle')
    arrow.set('w', 'sm')
    arrow.set('len', 'sm')
    line.append(arrow)
    return shape


def metric_card(slide, label, value, note, x, y, w, h, *, accent=BLUE):
    card = rounded_box(slide, x, y, w, h, fill=PALE_2, line="D5E0EA")
    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h)
    )
    set_fill(stripe, accent)
    stripe.line.fill.background()
    compact = h < 1.3
    text_box(slide, label, x + 0.18, y + 0.1, w - 0.32, 0.43, size=12, color=GRAY, bold=True)
    text_box(slide, value, x + 0.18, y + (0.52 if compact else 0.65), w - 0.32, 0.5, size=20 if compact else 24, color=accent, bold=True)
    if note and not compact:
        text_box(slide, note, x + 0.18, y + h - 0.32, w - 0.32, 0.22, size=10.5, color=GRAY)
    return card


def process_box(slide, title, subtitle, x, y, w, h, *, fill=PALE, line=BLUE):
    rounded_box(slide, x, y, w, h, fill=fill, line=line)
    text_box(slide, title, x + 0.1, y + 0.12, w - 0.2, 0.47, size=14, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, subtitle, x + 0.1, y + 0.62, w - 0.2, max(0.34,h - 0.68), size=11, color=GRAY, align=PP_ALIGN.CENTER)


def read_metrics() -> tuple[dict[str, dict[str, str]], dict[str, dict[str, dict[str, str]]]]:
    with (ROOT / "results/stage_smoke_tests/summary.csv").open(newline="", encoding="utf-8") as handle:
        summary = {row["stage"]: row for row in csv.DictReader(handle)}
    priority: dict[str, dict[str, dict[str, str]]] = {"stage1": {}, "stage2": {}}
    with (ROOT / "results/stage_smoke_tests/priority_stage_metrics.csv").open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            priority[row["stage"]][row["variable"]] = row
    return summary, priority


def add_bar(slide, label, value, max_value, x, y, w, *, color=BLUE, value_text=None):
    text_box(slide, label, x, y - 0.02, 1.35, 0.25, size=10.5, color=NAVY, bold=True)
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 1.4), Inches(y), Inches(w), Inches(0.22))
    set_fill(bg, "E2EAF1")
    bg.line.fill.background()
    fg = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x + 1.4),
        Inches(y),
        Inches(max(0.06, w * value / max_value)),
        Inches(0.22),
    )
    set_fill(fg, color)
    fg.line.fill.background()
    text_box(slide, value_text or f"{value:.3f}", x + 1.45 + w, y - 0.04, 0.8, 0.25, size=10, color=color, bold=True, align=PP_ALIGN.RIGHT)


def build() -> Path:
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"Missing template: {TEMPLATE}")
    summary, priority = read_metrics()
    prs = Presentation(str(TEMPLATE))
    clear_slides(prs)
    blank = prs.slide_layouts[6]

    # 1 — Cover
    slide = prs.slides.add_slide(blank)
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.28), Inches(0), Inches(0.11), Inches(7.5))
    set_fill(accent, BLUE)
    accent.line.fill.background()
    text_box(slide, "TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT TP. HỒ CHÍ MINH", 0.72, 0.55, 8.6, 0.28, size=11, color=GRAY, bold=True)
    text_box(slide, "NEURAL ATMOSPHERE OPERATOR", 0.72, 1.55, 11.3, 0.62, size=31, color=BLUE, bold=True)
    text_box(slide, "Dự báo khí quyển toàn cầu bằng Spherical Fourier Neural Operator", 0.72, 2.25, 11.0, 0.55, size=21, color=NAVY, bold=True)
    rounded_box(slide, 0.72, 3.12, 4.55, 0.52, fill=PALE, line="C9D9E7")
    text_box(slide, "STAGE 1  ·  BASELINE SFNO", 0.87, 3.27, 4.2, 0.22, size=12.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    rounded_box(slide, 5.48, 3.12, 5.35, 0.52, fill=GREEN_PALE, line="B8D8C5")
    text_box(slide, "STAGE 2  ·  SPHERICAL RESIDUAL CORRECTOR", 5.63, 3.27, 5.05, 0.22, size=12.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "Sinh viên: Lương Nguyễn Thành Hưng  ·  MSSV: 23110111", 0.76, 5.55, 8.5, 0.3, size=13, color=NAVY)
    text_box(slide, "TP. Hồ Chí Minh, 2026", 0.76, 6.03, 4.0, 0.24, size=10.5, color=GRAY)
    text_box(slide, "01", 12.16, 6.82, 0.52, 0.25, size=9, color=BLUE, bold=True, align=PP_ALIGN.RIGHT)

    # 2 — Problem and goals
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Bài toán và mục tiêu", 2)
    rounded_box(slide, 0.48, 1.14, 4.0, 5.55, fill=PALE, line="CBD8E4")
    text_box(slide, "Bài toán", 0.78, 1.45, 3.4, 0.34, size=20, color=BLUE, bold=True)
    add_bullets(slide, [
        "Dự báo trạng thái khí quyển toàn cầu sau 6 giờ",
        "71 kênh ERA5 trên lưới 0,5° (361 × 720)",
        "Học toán tử tiến hóa thay vì mô hình tại từng điểm",
        "Giới hạn thực tế: bộ nhớ, thời gian và ngân sách GPU",
    ], 0.75, 2.0, 3.35, 3.5, size=15.5, spacing=13)
    text_box(slide, "Câu hỏi nghiên cứu", 4.95, 1.23, 4.2, 0.38, size=21, color=NAVY, bold=True)
    questions = [
        ("Q1", "SFNO baseline học được bao nhiêu từ một trạng thái hiện tại?"),
        ("Q2", "History + forcing + residual corrector có giảm sai số không?"),
        ("Q3", "Đổi lại, mô hình mất gì khi cắt mạnh dao động tần số cao?"),
    ]
    for i, (tag, question) in enumerate(questions):
        y = 1.83 + i * 1.35
        rounded_box(slide, 4.95, y, 7.55, 1.05, fill=WHITE, line="CAD8E5")
        rounded_box(slide, 5.13, y + 0.2, 0.72, 0.62, fill=BLUE if i < 2 else RED, line=BLUE if i < 2 else RED)
        text_box(slide, tag, 5.18, y + 0.37, 0.62, 0.2, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, question, 6.08, y + 0.23, 6.05, 0.58, size=15, color=NAVY, bold=i == 2, valign=MSO_ANCHOR.MIDDLE)
    rounded_box(slide, 4.95, 6.02, 7.55, 0.65, fill=GREEN_PALE, line="B8D8C5")
    text_box(slide, "Phạm vi báo cáo: chỉ Stage 1 và Stage 2", 5.18, 6.22, 7.05, 0.22, size=13, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 3 — Data
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Dữ liệu và giao thức đánh giá", 3, "Tách thời gian tuyệt đối; thống kê chỉ được tính trên tập train")
    text_box(slide, "ERA5 / WeatherBench2", 0.53, 1.17, 3.1, 0.3, size=18, color=BLUE, bold=True)
    stages = [
        ("TRAIN", "1983–2018", "36 năm", BLUE, PALE),
        ("VALIDATION", "2019", "Chọn checkpoint", GREEN, GREEN_PALE),
        ("TEST", "07/2020", "Held-out", RED, RED_PALE),
    ]
    starts = [0.55, 4.52, 8.49]
    widths = [3.45, 3.45, 3.95]
    for x, w, (name, years, note, color, fill) in zip(starts, widths, stages):
        rounded_box(slide, x, 1.72, w, 1.35, fill=fill, line=color)
        text_box(slide, name, x + 0.15, 1.92, w - 0.3, 0.25, size=11, color=color, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, years, x + 0.15, 2.25, w - 0.3, 0.35, size=21, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, note, x + 0.15, 2.68, w - 0.3, 0.2, size=9.5, color=GRAY, align=PP_ALIGN.CENTER)
    connector(slide, 4.04, 2.4, 4.45, 2.4, MID, 1.5)
    connector(slide, 8.01, 2.4, 8.42, 2.4, MID, 1.5)
    cards = [
        ("71", "kênh khí quyển", "6 surface + 5 biến × 13 mức"),
        ("6 h", "cadence", "dự báo one-step K=1"),
        ("0,5°", "độ phân giải", "regrid bảo toàn từ 0,25°"),
        ("95", "khởi tạo tối đa", "cửa sổ 7 ngày trong 07/2020"),
    ]
    for i, (value, label, note) in enumerate(cards):
        metric_card(slide, label, value, note, 0.55 + i * 3.1, 3.58, 2.75, 1.5, accent=BLUE if i < 3 else GREEN)
    rounded_box(slide, 0.55, 5.48, 12.0, 1.08, fill=PALE_2, line="D6E1EB")
    text_box(slide, "Tiền xử lý", 0.8, 5.72, 1.35, 0.25, size=14, color=BLUE, bold=True)
    text_box(slide, "Chuẩn hóa theo σ(x)  →  dự báo standardized tendency  →  hoàn nguyên về đơn vị vật lý", 2.05, 5.68, 9.9, 0.3, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "Latitude-cell-area weighting giảm thiên lệch do mật độ điểm lưới tại hai cực.", 2.05, 6.06, 9.9, 0.22, size=10.5, color=GRAY, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 4 — Overall methodology and loss
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Phương pháp tổng thể", 4, "Cùng một objective để so sánh công bằng giữa hai stage")
    flow = [
        ("Trạng thái", "x(t)", 0.48, PALE),
        ("SFNO", "không gian phổ cầu", 2.75, PALE),
        ("Tendency", "Δx / σ(Δx)", 5.32, PALE),
        ("Residual update", "x(t)+s·Δx", 7.95, GREEN_PALE),
        ("Dự báo", "x(t+6h)", 10.65, GREEN_PALE),
    ]
    for title, subtitle, x, fill in flow:
        process_box(slide, title, subtitle, x, 1.45, 2.0, 1.05, fill=fill, line=GREEN if fill == GREEN_PALE else BLUE)
    for x1, x2 in [(2.49, 2.69), (4.76, 5.26), (7.34, 7.89), (9.96, 10.59)]:
        connector(slide, x1, 1.98, x2, 1.98)
    rounded_box(slide, 0.48, 3.05, 7.72, 3.5, fill=PALE_2, line="D1DDE8")
    text_box(slide, "Loss huấn luyện", 0.78, 3.34, 2.7, 0.35, size=20, color=BLUE, bold=True)
    text_box(slide, "L = Σc wc · Earea [ ((ŷc − yc) / sc)² ]", 0.82, 4.02, 7.0, 0.55, size=25, color=NAVY, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_bullets(slide, [
        "sc = σ(Δxc) / σ(xc): cân bằng độ khó giữa các biến",
        "wc: trọng số biến và mức áp suất kiểu GraphCast",
        "Earea: trọng số diện tích ô lưới theo vĩ độ",
    ], 0.85, 4.85, 6.9, 1.3, size=13.5, spacing=6)
    rounded_box(slide, 8.55, 3.05, 4.0, 3.5, fill=RED_PALE, line="EDBDBD")
    text_box(slide, "Trade-off", 8.88, 3.36, 3.35, 0.34, size=20, color=RED, bold=True)
    add_bullets(slide, [
        "MSE ổn định và dễ tối ưu",
        "Nhưng ưu tiên sai số trung bình",
        "Cực trị hiếm dễ bị làm trơn",
        "Cắt phổ cao tạo sàn sai số",
    ], 8.85, 4.0, 3.15, 1.7, size=14.5, spacing=8)
    text_box(slide, "Loss giảm ≠ cực trị tốt hơn", 8.92, 5.97, 3.18, 0.3, size=13.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 5 — Stage 1
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Stage 1 — SFNO baseline", 5, "Một trạng thái hiện tại → standardized tendency sau 6 giờ")
    process_box(slide, "x(t)", "71 kênh", 0.62, 1.55, 1.55, 1.0)
    process_box(slide, "Encoder", "71 → 192", 2.75, 1.55, 1.75, 1.0)
    process_box(slide, "6 SFNO blocks", "SHT → spectral mixing → iSHT", 5.08, 1.42, 3.05, 1.28, fill=PALE, line=BLUE)
    process_box(slide, "Decoder", "192 → 71", 8.72, 1.55, 1.75, 1.0)
    process_box(slide, "x(t+6h)", "residual output", 11.02, 1.55, 1.65, 1.0, fill=GREEN_PALE, line=GREEN)
    for x1, x2 in [(2.2, 2.69), (4.53, 5.02), (8.16, 8.66), (10.5, 10.96)]:
        connector(slide, x1, 2.05, x2, 2.05)
    specs = [
        ("E192 / L6", "Embedding / blocks"),
        ("Scale ×2", "Giảm resolution nội bộ"),
        ("K = 1", "Lead 6 giờ"),
        ("~25 epochs", "Best tại epoch 24"),
    ]
    for i, (value, note) in enumerate(specs):
        metric_card(slide, note, value, "", 0.62 + i * 3.03, 3.25, 2.68, 1.38, accent=BLUE)
    rounded_box(slide, 0.62, 5.05, 12.05, 1.42, fill=PALE_2, line="D5E0EA")
    text_box(slide, "Điểm mạnh", 0.88, 5.35, 1.35, 0.25, size=14, color=GREEN, bold=True)
    text_box(slide, "Ổn định · học động lực quy mô lớn · vượt persistence rõ rệt", 2.15, 5.31, 4.2, 0.35, size=13.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "Nút thắt", 6.65, 5.35, 1.25, 0.25, size=14, color=RED, bold=True)
    text_box(slide, "Thiếu history/forcing · phổ cao bị suy giảm · cực trị bị làm trơn", 7.83, 5.31, 4.45, 0.35, size=13.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 6 — Stage 1 result
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Kết quả Stage 1 và điểm bão hòa", 6, "Smoke test: 12 khởi tạo ERA5, lead 6 giờ")
    s1 = summary["stage1"]
    metric_card(slide, "Best validation loss", "0,13444", "epoch 24 · 187.850 updates", 0.55, 1.25, 3.0, 1.48, accent=BLUE)
    metric_card(slide, "Normalized RMSE", f"{float(s1['normalized_rmse']):.4f}".replace(".", ","), "persistence: 0,3312", 3.82, 1.25, 2.85, 1.48, accent=BLUE)
    metric_card(slide, "Normalized ACC", f"{float(s1['normalized_acc']):.4f}".replace(".", ","), "persistence: 0,9233", 6.95, 1.25, 2.85, 1.48, accent=GREEN)
    metric_card(slide, "Objective loss", f"{float(s1['objective_loss']):.4f}".replace(".", ","), "standardized tendency MSE", 10.08, 1.25, 2.45, 1.48, accent=BLUE)
    rounded_box(slide, 0.55, 3.12, 7.22, 3.42, fill=PALE_2, line="D5E0EA")
    text_box(slide, "Model so với baseline", 0.83, 3.4, 3.4, 0.3, size=19, color=NAVY, bold=True)
    add_bar(slide, "Persistence", float(s1["persistence_rmse"]), 0.35, 0.9, 4.05, 5.2, color=MID)
    add_bar(slide, "Stage 1", float(s1["normalized_rmse"]), 0.35, 0.9, 4.75, 5.2, color=BLUE)
    text_box(slide, "RMSE giảm 52,5% so với persistence", 1.0, 5.45, 5.9, 0.35, size=16, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    rounded_box(slide, 8.08, 3.12, 4.45, 3.42, fill=RED_PALE, line="EDBDBD")
    text_box(slide, "Vì sao loss bão hòa?", 8.42, 3.42, 3.8, 0.3, size=19, color=RED, bold=True)
    add_bullets(slide, [
        "Input chỉ có một thời điểm",
        "Không có forcing địa hình/thời gian",
        "Scale factor làm nghèo dải tần cao",
        "MSE khuyến khích dự báo trơn",
    ], 8.37, 4.05, 3.55, 1.8, size=14.5, spacing=8)
    text_box(slide, "Capacity tăng thêm không phục hồi được thông tin đã mất.", 8.45, 5.93, 3.62, 0.42, size=11.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 7 — Stage 2 architecture
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Stage 2 — Spherical residual corrector", 7, "Giữ nguyên backbone Stage 1; chỉ học phần sai số còn lại")
    process_box(slide, "Stage 1 backbone", "SFNO E192/L6\nkhởi tạo từ best checkpoint", 0.55, 1.5, 3.05, 1.28, fill=PALE, line=BLUE)
    process_box(slide, "Baseline tendency", "71 kênh · detach", 4.18, 1.5, 2.35, 1.28, fill=PALE, line=BLUE)
    process_box(slide, "Cộng residual", "baseline + correction", 10.25, 1.5, 2.25, 1.28, fill=GREEN_PALE, line=GREEN)
    connector(slide, 3.65, 2.14, 4.12, 2.14)
    connector(slide, 6.58, 2.14, 10.18, 2.14)
    rounded_box(slide, 0.55, 3.35, 6.25, 2.92, fill=PALE_2, line="D4DFE9")
    text_box(slide, "Covariates cho corrector", 0.85, 3.66, 3.2, 0.32, size=19, color=NAVY, bold=True)
    inputs = [
        ("x(t)", "71"),
        ("Δx gần nhất", "71"),
        ("baseline", "71"),
        ("forcing", "7"),
        ("latitude", "2"),
    ]
    for i, (name, count) in enumerate(inputs):
        x = 0.85 + (i % 3) * 1.88
        y = 4.25 + (i // 3) * 0.82
        rounded_box(slide, x, y, 1.62, 0.58, fill=WHITE, line="C8D7E4")
        text_box(slide, name, x + 0.06, y + 0.11, 1.05, 0.2, size=10.5, color=NAVY, bold=True)
        text_box(slide, count, x + 1.06, y + 0.1, 0.4, 0.22, size=12, color=BLUE, bold=True, align=PP_ALIGN.RIGHT)
    text_box(slide, "Tổng: 222 kênh đầu vào", 4.55, 5.38, 1.85, 0.32, size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    rounded_box(slide, 7.15, 3.35, 5.35, 2.92, fill=GREEN_PALE, line="B8D8C5")
    text_box(slide, "Corrector SFNO", 7.48, 3.68, 2.5, 0.32, size=20, color=GREEN, bold=True)
    add_bullets(slide, [
        "Embedding 84 · 4 spectral blocks",
        "Native grid 361 × 720 · scale factor 1",
        "Không pooling / bilinear upsampling",
        "Projection cuối khởi tạo bằng zero",
    ], 7.48, 4.2, 4.35, 1.6, size=14, spacing=7)
    connector(slide, 6.82, 4.83, 7.08, 4.83, GREEN, 2.0)
    connector(slide, 9.82, 3.3, 10.95, 2.84, GREEN, 2.0)
    add_footer(slide)

    # 8 — Conditioning and fine-tuning
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Conditioning và progressive fine-tuning", 8)
    rounded_box(slide, 0.52, 1.14, 5.05, 5.55, fill=PALE, line="CBD9E5")
    text_box(slide, "7 forcing channels", 0.86, 1.48, 3.0, 0.35, size=20, color=BLUE, bold=True)
    add_bullets(slide, [
        "Orography chuẩn hóa",
        "Land–sea mask",
        "Cosine solar zenith angle",
        "Sin / cos local time",
        "Sin / cos annual progress",
    ], 0.88, 2.05, 4.05, 2.3, size=15, spacing=9)
    rounded_box(slide, 0.88, 4.75, 4.25, 1.3, fill=WHITE, line="C8D7E4")
    text_box(slide, "Mục đích", 1.1, 5.0, 1.0, 0.25, size=12, color=GREEN, bold=True)
    text_box(slide, "Cho model biết địa hình, mùa và chu kỳ ngày–đêm mà state tức thời không biểu diễn trực tiếp.", 2.02, 4.92, 2.78, 0.65, size=12.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text_box(slide, "Lịch fine-tune", 6.05, 1.33, 2.7, 0.35, size=21, color=NAVY, bold=True)
    # Timeline
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(6.13), Inches(2.28), Inches(5.86), Inches(0.1))
    set_fill(bar, MID)
    bar.line.fill.background()
    for x, label, sub, color in [
        (6.2, "Epoch 0–1", "Corrector only", BLUE),
        (9.15, "Epoch 2+", "Corrector + tail", GREEN),
    ]:
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(2.03), Inches(0.55), Inches(0.55))
        set_fill(circle, color)
        circle.line.fill.background()
        text_box(slide, label, x - 0.28, 2.76, 1.25, 0.25, size=12, color=color, bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, sub, x - 0.46, 3.11, 1.65, 0.24, size=11, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    rounded_box(slide, 6.05, 3.7, 6.42, 2.55, fill=PALE_2, line="D4DFE9")
    rows = [
        ("Corrector LR", "1 × 10⁻⁴", BLUE),
        ("Backbone tail LR", "1 × 10⁻⁵", GREEN),
        ("Unfreeze", "block cuối + decoder", NAVY),
        ("Gradient clip", "1,0", NAVY),
    ]
    for i, (label, value, color) in enumerate(rows):
        y = 4.02 + i * 0.48
        text_box(slide, label, 6.35, y, 2.45, 0.22, size=12, color=GRAY, bold=True)
        text_box(slide, value, 8.8, y, 3.15, 0.22, size=13, color=color, bold=True, align=PP_ALIGN.RIGHT)
    text_box(slide, "Backbone LR nhỏ hơn 10× để hạn chế catastrophic forgetting.", 6.4, 5.84, 5.6, 0.24, size=10.5, color=GRAY, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 9 — Main comparison
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Stage 1 vs Stage 2", 9, "Checkpoint được chọn bằng validation; bảng dưới là smoke test 6 giờ")
    s2 = summary["stage2"]
    val1, val2 = 0.13444305187815098, 0.12538444922294145
    metrics = [
        ("Best validation loss", val1, val2, "−6,74%", 0.15),
        ("Objective loss", float(s1["objective_loss"]), float(s2["objective_loss"]), "−3,49%", 0.16),
        ("Normalized RMSE", float(s1["normalized_rmse"]), float(s2["normalized_rmse"]), "−1,48%", 0.18),
    ]
    for i, (name, a, b, gain, vmax) in enumerate(metrics):
        y = 1.35 + i * 1.55
        rounded_box(slide, 0.55, y, 8.25, 1.25, fill=PALE_2, line="D7E1EA")
        text_box(slide, name, 0.83, y + 0.19, 2.2, 0.26, size=14, color=NAVY, bold=True)
        add_bar(slide, "Stage 1", a, vmax, 3.0, y + 0.22, 3.5, color=MID, value_text=f"{a:.4f}")
        add_bar(slide, "Stage 2", b, vmax, 3.0, y + 0.68, 3.5, color=GREEN, value_text=f"{b:.4f}")
        rounded_box(slide, 7.62, y + 0.37, 0.92, 0.52, fill=GREEN_PALE, line="B8D8C5")
        text_box(slide, gain, 7.7, y + 0.53, 0.75, 0.2, size=12, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    rounded_box(slide, 9.12, 1.35, 3.4, 4.35, fill=GREEN_PALE, line="B8D8C5")
    text_box(slide, "ACC", 9.48, 1.76, 2.7, 0.28, size=17, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "0,983585", 9.48, 2.27, 2.7, 0.42, size=22, color=GRAY, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "↓ Stage 1", 9.48, 2.72, 2.7, 0.22, size=10, color=GRAY, align=PP_ALIGN.CENTER)
    text_box(slide, "0,984121", 9.48, 3.32, 2.7, 0.42, size=25, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "Stage 2", 9.48, 3.79, 2.7, 0.22, size=11, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "+0,000536", 9.48, 4.47, 2.7, 0.35, size=17, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "tăng tuyệt đối", 9.48, 4.9, 2.7, 0.22, size=10, color=GRAY, align=PP_ALIGN.CENTER)
    rounded_box(slide, 0.55, 6.15, 11.97, 0.48, fill=PALE, line="CAD8E5")
    text_box(slide, "Kết luận: Stage 2 cải thiện nhất quán nhưng biên lợi ích nhỏ — phù hợp với một residual corrector, không phải bước nhảy kiến trúc.", 0.78, 6.29, 11.5, 0.2, size=12.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 10 — Physical variables
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Cải thiện theo biến vật lý", 10, "Phần trăm giảm RMSE của Stage 2 so với Stage 1 tại lead 6 giờ")
    variables = ["Z500", "T850", "T2M", "U10", "V10", "MSLP", "Q850", "U850", "V850"]
    gains = []
    for variable in variables:
        a = float(priority["stage1"][variable]["rmse"])
        b = float(priority["stage2"][variable]["rmse"])
        gains.append((variable, 100.0 * (a - b) / a, a, b, priority["stage2"][variable]["unit"]))
    max_gain = max(item[1] for item in gains) * 1.12
    for i, (variable, gain, _a, _b, _unit) in enumerate(gains):
        col, row = i // 5, i % 5
        x = 0.62 + col * 6.25
        y = 1.35 + row * 0.89
        add_bar(slide, variable, gain, max_gain, x, y, 3.72, color=GREEN, value_text=f"{gain:.1f}%")
    rounded_box(slide, 0.62, 5.93, 12.0, 0.72, fill=PALE_2, line="D5E0EA")
    text_box(slide, "Giảm mạnh nhất", 0.9, 6.15, 1.38, 0.22, size=11, color=GRAY, bold=True)
    best = max(gains, key=lambda item: item[1])
    text_box(slide, f"{best[0]}  {best[2]:.4f} → {best[3]:.4f} {best[4]}  ({best[1]:.1f}%)", 2.18, 6.09, 3.5, 0.3, size=14, color=GREEN, bold=True)
    text_box(slide, "Tất cả 9 biến ưu tiên đều giảm RMSE", 7.05, 6.12, 4.85, 0.25, size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 11 — Interpretation and limitations
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Diễn giải kết quả và giới hạn", 11)
    rounded_box(slide, 0.52, 1.15, 5.9, 5.55, fill=GREEN_PALE, line="B8D8C5")
    text_box(slide, "Vì sao Stage 2 tốt hơn?", 0.88, 1.5, 4.8, 0.35, size=21, color=GREEN, bold=True)
    add_bullets(slide, [
        "Học residual nhỏ thay vì học lại toàn bộ dự báo",
        "History cung cấp hướng chuyển động gần nhất",
        "Forcing bổ sung địa hình và chu kỳ thời gian",
        "Corrector full-resolution giữ thêm chi tiết phổ",
        "Progressive unfreezing bảo vệ backbone đã học",
    ], 0.88, 2.05, 4.95, 2.7, size=15, spacing=10)
    rounded_box(slide, 0.9, 5.2, 5.12, 0.92, fill=WHITE, line="C4D9CC")
    text_box(slide, "Yếu tố có khả năng đóng góp lớn nhất", 1.08, 5.4, 4.7, 0.2, size=11, color=GRAY, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "Residual corrector + transfer learning", 1.08, 5.72, 4.7, 0.24, size=14, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    rounded_box(slide, 6.72, 1.15, 5.9, 5.55, fill=RED_PALE, line="EDBDBD")
    text_box(slide, "Những gì chưa được giải quyết", 7.08, 1.5, 4.9, 0.35, size=21, color=RED, bold=True)
    add_bullets(slide, [
        "Mới huấn luyện và kiểm tra chính ở K=1 / 6 giờ",
        "Smoke test chỉ có 12 mẫu; chưa đủ kiểm định thống kê",
        "Chưa có ablation để tách đóng góp từng thành phần",
        "MSE vẫn làm nhẹ ảnh hưởng của sự kiện cực đoan",
        "Thông tin phổ cao đã mất không thể tái tạo hoàn toàn",
    ], 7.08, 2.05, 4.95, 2.7, size=15, spacing=10)
    rounded_box(slide, 7.1, 5.2, 5.12, 0.92, fill=WHITE, line="E0C7C7")
    text_box(slide, "Thông điệp quan trọng", 7.28, 5.4, 4.7, 0.2, size=11, color=GRAY, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "Loss giảm không đảm bảo cực trị tốt hơn", 7.28, 5.72, 4.7, 0.24, size=14, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 12 — Conclusion
    slide = prs.slides.add_slide(blank)
    add_title(slide, "Kết luận và hướng phát triển", 12)
    rounded_box(slide, 0.55, 1.2, 5.95, 4.85, fill=PALE, line="CBD9E5")
    text_box(slide, "Kết luận", 0.9, 1.56, 2.1, 0.34, size=22, color=BLUE, bold=True)
    add_bullets(slide, [
        "Stage 1 tạo baseline SFNO ổn định cho dự báo toàn cầu.",
        "Stage 2 giảm validation loss 6,74% bằng residual correction.",
        "History, forcing và fine-tuning có kiểm soát tạo cải thiện nhất quán.",
        "Giới hạn lớn nhất là loss trung bình và suy giảm dải tần cao.",
    ], 0.9, 2.15, 5.0, 2.65, size=15.5, spacing=12)
    text_box(slide, "Đóng góp chính: một pipeline hai stage có checkpoint, evaluator và giao thức tái lập rõ ràng.", 0.98, 5.15, 4.95, 0.56, size=13.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    rounded_box(slide, 6.8, 1.2, 5.8, 4.85, fill=GREEN_PALE, line="B8D8C5")
    text_box(slide, "Hướng phát triển", 7.15, 1.56, 2.8, 0.34, size=22, color=GREEN, bold=True)
    add_bullets(slide, [
        "Ablation theo số mode và resolution",
        "Loss phổ / gradient / extreme-aware",
        "Sampling ưu tiên bão, front và cực trị",
        "Multi-step curriculum và đánh giá 3–7 ngày",
        "Nhánh đa phân giải để giữ cấu trúc tần số cao",
    ], 7.15, 2.15, 4.75, 2.65, size=15.5, spacing=11)
    rounded_box(slide, 7.15, 5.12, 4.95, 0.62, fill=WHITE, line="C5DBCE")
    text_box(slide, "Mục tiêu tiếp theo: tối ưu chất lượng cực trị, không chỉ tối ưu average MSE.", 7.35, 5.28, 4.55, 0.28, size=12.5, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    text_box(slide, "CẢM ƠN THẦY/CÔ VÀ CÁC BẠN ĐÃ THEO DÕI", 1.1, 6.43, 11.1, 0.35, size=19, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    print(build())

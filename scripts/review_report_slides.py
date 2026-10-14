"""Evidence plots and per-slide editorial repairs for the report."""

from __future__ import annotations

import csv
import hashlib
import shutil
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from report_wording import apply_wording, correct_figure_terms

from build_report_slides import (
    BLUE, GREEN, GRAY, MID, NAVY, PALE, RED, WHITE,
    add_bullets, connector, process_box, rounded_box, text_box, rgb,
)


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / 'reports' / 'evidence'
FIGURES = ROOT / 'reports' / 'figures'


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding='utf-8', newline='') as stream:
        return list(csv.DictReader(stream))


def evidence_histories():
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    sources = {}
    for stage, folder in [(1, 'stage1_baseline'), (2, 'stage2_spherical_corrector')]:
        target = EVIDENCE / f'stage{stage}_history.csv'
        source = ROOT.parent / 'NeuralAtmosphereOperator_COMPLETE_EXPORT' / '01_CHECKPOINTS' / folder / 'history.csv'
        if source.is_file():
            shutil.copyfile(source, target)
        if not target.is_file():
            raise FileNotFoundError(f'Missing measured training history: {target}')
        rows = read_csv(target)
        sources[stage] = rows
        expected = 0.13444305187815098 if stage == 1 else 0.12538444922294145
        assert abs(min(float(row['validation_loss']) for row in rows)-expected) < 1e-10
    return sources


def axes_style(ax, title, xlabel='Epoch (1-based)', ylabel='Loss'):
    ax.set_title(title, loc='left', fontsize=15, fontweight='bold', color='#20384F', pad=16)
    ax.set_xlabel(xlabel, fontsize=12); ax.set_ylabel(ylabel, fontsize=12)
    ax.spines[['top','right']].set_visible(False)
    ax.spines[['left','bottom']].set_color('#A9B7C6')
    ax.tick_params(labelsize=11, colors='#455564')
    ax.grid(color='#E4EBF2', alpha=.8); ax.set_axisbelow(True)


def save(fig, name):
    correct_figure_terms(fig)
    path = FIGURES / f'{name}.png'
    fig.savefig(path, dpi=240, bbox_inches='tight', facecolor='white')
    fig.savefig(path.with_suffix('.pdf'), bbox_inches='tight', facecolor='white')
    plt.close(fig)
    return path


def measured_plots(histories):
    FIGURES.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({'font.family':'DejaVu Sans','mathtext.fontset':'stix', 'font.size':12})
    paths = {}
    for stage, rows in histories.items():
        x=np.array([int(row['epoch']) for row in rows])
        train=np.array([float(row['train_loss']) for row in rows])
        valid=np.array([float(row['validation_loss']) for row in rows])
        fig, ax=plt.subplots(figsize=(10,4.65),layout='constrained')
        ax.plot(x,train,color='#1976D2',lw=2.6,marker='o',ms=4,label='Train')
        ax.plot(x,valid,color='#2E8B57',lw=2.6,marker='o',ms=4,label='Validation')
        if stage==2:
            ax.axvspan(.8,2.5,color='#1976D2',alpha=.055)
            ax.axvline(2.5,color='#D94C4C',ls='--',lw=1.3,label='Unfreeze tail')
            ax.set_xticks(x); ax.set_xlim(.8,4.3)
            ax.axhline(.13444305187815098,color='#7D8B99',ls=':',lw=1.7,label='Stage 1 best validation')
        best=int(np.argmin(valid))
        ax.scatter(x[best],valid[best],s=100,facecolor='white',edgecolor='#2E8B57',linewidth=2,zorder=5)
        ax.annotate(f'Best {valid[best]:.6f}',(x[best],valid[best]),xytext=(-110,26),textcoords='offset points',fontsize=11,color='#19663D',arrowprops={'arrowstyle':'->','color':'#19663D'})
        axes_style(ax,f'Stage {stage}: train / validation loss')
        ax.legend(frameon=False,fontsize=10,loc='upper right' if stage==1 else 'center',
                  bbox_to_anchor=None if stage==1 else (.47,.42), ncol=1)
        paths[f'loss{stage}']=save(fig,f'stage{stage}_loss')

    rows=histories[1]; tail=rows[-10:]
    x=np.array([int(r['epoch']) for r in tail]); valid=np.array([float(r['validation_loss']) for r in tail])
    fig,axes=plt.subplots(1,2,figsize=(11.3,4.4),layout='constrained')
    axes[0].plot(x,valid,'o-',color='#2E8B57',lw=2.5)
    axes_style(axes[0],'Stage 1: 10 epoch cuối')
    improvements=100*(valid[:-1]-valid[1:])/valid[:-1]
    axes[1].bar(x[1:],improvements,color='#1976D2')
    axes_style(axes[1],'Mức giảm validation từng epoch',ylabel='Reduction (%)')
    paths['convergence']=save(fig,'stage1_convergence')

    fig,axes=plt.subplots(1,2,figsize=(11.3,4.45),layout='constrained')
    for i,stage in enumerate((1,2)):
        rows=histories[stage]; x=[int(r['epoch']) for r in rows]
        axes[i].plot(x,[float(r['learning_rate']) for r in rows],'o-',color='#1976D2',lw=2,label='Model / corrector')
        if stage==2:
            # The saved epoch-end LR can be nonzero while requires_grad remains false.
            axes[i].plot(x,[float(r['backbone_learning_rate']) if r['fine_tuning_phase']!='corrector_only' else 0 for r in rows],'s-',color='#2E8B57',lw=2,label='Active backbone tail')
            axes[i].axvline(2.5,color='#D94C4C',ls='--'); axes[i].set_xticks(x)
        axes_style(axes[i],f'Stage {stage}: learning rate',ylabel='Epoch-end LR'); axes[i].ticklabel_format(axis='y',style='sci',scilimits=(0,0)); axes[i].legend(frameon=False,fontsize=10)
    paths['learning_rates']=save(fig,'learning_rates')

    fig,axes=plt.subplots(1,2,figsize=(11.3,4.4),layout='constrained')
    for i,stage in enumerate((1,2)):
        rows=histories[stage]; x=[int(r['epoch']) for r in rows]
        axes[i].plot(x,[float(r['mean_preclip_gradient_norm']) for r in rows],'o-',color='#1976D2',label='Mean before clipping')
        axes[i].plot(x,[float(r['max_preclip_gradient_norm']) for r in rows],'s--',color='#D94C4C',label='Max before clipping')
        axes[i].axhline(1,color='#7D8B99',ls=':',label='Clip threshold'); axes[i].set_yscale('log')
        axes_style(axes[i],f'Stage {stage}: gradient stability',ylabel='Gradient norm (log)'); axes[i].legend(frameon=False,fontsize=9)
    paths['gradients']=save(fig,'gradient_stability')

    rows=read_csv(ROOT/'results/stage_smoke_tests/priority_stage_metrics.csv')
    by={(r['stage'],r['variable']):r for r in rows}
    variables=['Z500','T850','T2M','U10','V10','MSLP','Q850','U850','V850']
    for metric in ['rmse','bias']:
        fig,axes=plt.subplots(3,3,figsize=(11.4,6.3),layout='constrained')
        for ax,name in zip(axes.flat,variables):
            a=float(by[('stage1',name)][metric]); b=float(by[('stage2',name)][metric])
            ax.bar(['S1','S2'],[a,b],color=['#A9B7C6','#2E8B57'],width=.5)
            ax.axhline(0,color='#7D8B99',lw=.8)
            for j,v in enumerate([a,b]):
                ax.annotate(f'{v:.4f}',(j,v),xytext=(0,4 if v>=0 else -13),textcoords='offset points',ha='center',fontsize=10)
            axes_style(ax,f"{name} · {by[('stage1',name)]['unit']}",xlabel='',ylabel='')
            ax.margins(y=.28); ax.tick_params(labelsize=10)
        paths[metric+'_physical']=save(fig,metric+'_physical')

    all_rows=read_csv(ROOT/'results/stage_smoke_tests/all_stage_metrics.csv')
    pairs={(r['stage'],r['variable']):r for r in all_rows}
    prefixes=['u_component_of_wind','v_component_of_wind','geopotential','temperature','specific_humidity']
    levels=sorted({int(r['variable'].split('@')[1].replace('hPa','')) for r in all_rows if '@' in r['variable']})
    values=[]
    for prefix in prefixes:
        band=[]
        for level in levels:
            name=f'{prefix}@{level}hPa'; a=float(pairs[('stage1',name)]['rmse']); b=float(pairs[('stage2',name)]['rmse']); band.append(100*(a-b)/a)
        values.append(band)
    fig,ax=plt.subplots(figsize=(11.3,4.3),layout='constrained')
    limit=max(abs(np.array(values).min()),abs(np.array(values).max()))
    image=ax.imshow(values,cmap='RdYlGn',vmin=-limit,vmax=limit,aspect='auto')
    ax.set_xticks(range(len(levels)),[str(n) for n in levels]); ax.set_yticks(range(5),['U wind','V wind','Z','Temperature','Humidity'])
    for i in range(5):
        for j in range(len(levels)):ax.text(j,i,f'{values[i][j]:+.1f}',ha='center',va='center',fontsize=10,color='#20384F')
    ax.set_xlabel('Pressure (hPa)'); ax.set_title('RMSE reduction (%) by variable / pressure level',loc='left',fontsize=15,fontweight='bold',pad=16)
    fig.colorbar(image,ax=ax,label='Positive = Stage 2 better',shrink=.85)
    paths['pressure']=save(fig,'pressure_level_heatmap')
    return paths


def remove_content(slide):
    for shape in list(slide.shapes):
        # Preserve top branding and footer, remove content only.
        if Inches(.95) < shape.top < Inches(7.0):
            shape._element.getparent().remove(shape._element)


def image_fit(slide,path,x,y,w,h):
    with Image.open(path) as source: ratio=source.width/source.height
    aw=min(w,h*ratio); ah=aw/ratio
    slide.shapes.add_picture(str(path),Inches(x+(w-aw)/2),Inches(y+(h-ah)/2),width=Inches(aw),height=Inches(ah))


def chart_slide(slide,path,caption):
    remove_content(slide)
    image_fit(slide,path,.6,1.1,12.1,5.35)
    text_box(slide,caption,.85,6.52,11.6,.4,size=12,color=GRAY,align=PP_ALIGN.CENTER)


def insert_before(prs,slide,index):
    ids=prs.slides._sldIdLst
    element=ids[-1]; ids.remove(element); ids.insert(index,element)


def repair_and_enrich(prs,logo,standard_slide):
    histories=evidence_histories(); paths=measured_plots(histories)
    initial=list(prs.slides)
    def title(slide):
        return next((shape.text for shape in slide.shapes if shape.has_text_frame and Inches(.1)<shape.top<Inches(.5) and shape.left>Inches(.9)), '')

    # Repair actual ordering bugs discovered by opening each of the 57 slides.
    slide=initial[11]; remove_content(slide)
    for i,(a,b) in enumerate([('ERA5 source','0,25°'),('Area overlap','conservative'),('Periodic longitude','0° / 360°'),('Target grid','0,5°')]):
        x=.6+i*3.1; process_box(slide,a,b,x,1.35,2.5,1.2)
        if i<3:connector(slide,x+2.52,1.95,x+3.02,1.95)
    rounded_box(slide,.7,3.15,5.8,2.62,fill=PALE,line=MID)
    rounded_box(slide,6.85,3.15,5.8,2.62,fill=PALE,line=MID)
    text_box(slide,'Kiểm tra dữ liệu',1.0,3.4,5.0,.36,size=20,color=BLUE,bold=True)
    add_bullets(slide,['Channel order cố định','Tọa độ nhất quán; finite values','Contract hash kiểm soát resume'],1.0,4.02,5.0,1.4,size=18,spacing=12)
    text_box(slide,'Đánh đổi resolution',7.15,3.4,5.0,.36,size=20,color=RED,bold=True)
    add_bullets(slide,['Giảm bộ nhớ và chi phí','Nyquist thấp hơn','Mất chi tiết dưới grid scale'],7.15,4.02,5.0,1.4,size=18,spacing=12)

    # Model diagrams must show branches rather than a fictitious sequential skip.
    slide=initial[18]; remove_content(slide)
    text_box(slide,'Spectral branch và residual/MLP là hai đường xử lý trong block.',.8,1.05,11.7,.45,size=18,color=NAVY,bold=True)
    coords=[(.65,'vₖ','features'),(3.0,'SHT','sphere → modes'),(5.35,'Spectral filter','learned weights'),(7.7,'iSHT','modes → sphere'),(10.05,'Add + activation','output features')]
    for x,a,b in coords:process_box(slide,a,b,x,2.35,2.0,1.2)
    for i in range(4):connector(slide,coords[i][0]+2.02,2.95,coords[i+1][0]-.05,2.95)
    connector(slide,1.65,2.3,1.65,1.85,GRAY);connector(slide,1.65,1.85,11.05,1.85,GRAY);connector(slide,11.05,1.85,11.05,2.3,GRAY)
    text_box(slide,'residual / local path',5.0,1.49,3.4,.3,size=14,color=GRAY,align=PP_ALIGN.CENTER)
    process_box(slide,'MLP + residual','feature refinement',5.35,4.75,2.0,1.12)
    connector(slide,11.05,3.6,11.05,4.35);connector(slide,11.05,4.35,6.35,4.35);connector(slide,6.35,4.35,6.35,4.7)
    text_box(slide,'Sơ đồ chức năng; thứ tự normalization/skip chi tiết theo torch-harmonics 0.7.4.',.85,6.32,11.7,.4,size=12,color=GRAY,align=PP_ALIGN.CENTER)

    slide=initial[37]; remove_content(slide)
    process_box(slide,'x(t)','71 channels',.6,1.3,2.05,1.12)
    process_box(slide,'Backbone E192/L6','từ Stage 1',3.25,1.3,2.7,1.12)
    process_box(slide,'Baseline tendency','71 channels',6.5,1.3,2.45,1.12)
    process_box(slide,'Add correction','71-channel tendency',10.15,1.3,2.55,1.12)
    for a,b in [(2.7,3.2),(6.0,6.45),(9.0,10.1)]:connector(slide,a,1.86,b,1.86)
    process_box(slide,'Recent tendency','[x(t)−x(t−6h)] / s',.6,3.1,2.9,1.12)
    process_box(slide,'Forcing + geometry','7 + 2 channels',.6,4.7,2.9,1.12)
    process_box(slide,'Concatenate','71 + 71 + 71 + 7 + 2',4.55,3.65,3.2,1.2)
    process_box(slide,'Corrector E84/L4','222 → 71 channels',8.9,3.65,3.8,1.2)
    connector(slide,3.55,3.66,4.5,4.0,GREEN);connector(slide,3.55,5.26,4.5,4.5,GREEN)
    connector(slide,1.62,2.48,1.62,2.8,GREEN);connector(slide,1.62,2.8,6.15,2.8,GREEN);connector(slide,6.15,2.8,6.15,3.6,GREEN)
    connector(slide,7.73,2.48,7.73,3.03,GREEN);connector(slide,7.73,3.03,6.8,3.03,GREEN);connector(slide,6.8,3.03,6.8,3.6,GREEN)
    text_box(slide,'stop-gradient',7.85,2.68,1.95,.24,size=12,color=GRAY)
    connector(slide,7.8,4.25,8.85,4.25,GREEN)
    connector(slide,11.3,3.6,11.3,2.48,GREEN)
    text_box(slide,'Input lịch sử có 142 channels; backbone nhận 71; corrector nhận 222.',1.0,6.43,11.2,.34,size=16,color=NAVY,bold=True,align=PP_ALIGN.CENTER)

    chart_slide(initial[29],paths['rmse_physical'],'RMSE vật lý tại 6 giờ · 12 khởi tạo/stage · mỗi panel có đơn vị riêng; trục độc lập.')
    chart_slide(initial[30],paths['convergence'],'Nguồn: stage1_history.csv, epoch 16–25. Validation vẫn giảm ở cuối run; chưa chứng minh plateau.')
    chart_slide(initial[47],paths['bias_physical'],'Bias vật lý tại 6 giờ · dấu biểu thị lệch hệ thống; mỗi panel có đơn vị riêng.')
    # Remove attribution unsupported by ablations and acknowledge actual group regressions.
    slide=initial[45]
    for shape in list(slide.shapes):
        if shape.has_text_frame and shape.text.startswith('•'):
            shape._element.getparent().remove(shape._element)
    add_bullets(slide,['Surface: +1,72%','Humidity: +2,49%','U/V và T: trung vị âm','Không đồng đều theo pressure','Chưa có ablation để quy nguyên nhân'],9.86,2.15,2.28,3.1,size=14.5,spacing=12)

    # Correct claims that were stronger than the available evidence.
    replacements={
        'Vì sao Stage 1 bị bão hòa?':'Giả thuyết về giới hạn của Stage 1',
        'Residual corrector là cải tiến hiệu quả nhất':'Stage 2 cải thiện loss của toàn hệ thống',
        'Thay thế một phần chi phí của mô hình số':'Nghiên cứu dự báo nhanh từ dữ liệu',
        'Cφ(·) = 0 tại update đầu':'Cφ(·) = 0 trước bước optimizer đầu',
        'Backbone tail LR':'Backbone tail LR',
        'scale factor = 2':'scale factor = 2 · internal grid 181 × 360',
        'Điểm mạnh của Stage 1':'Stage 1 — hội tụ ở 10 epoch cuối',
        'Stage 1 trên các biến ưu tiên':'RMSE vật lý — Stage 1 và Stage 2',
        'SFNO baseline · kết quả · bão hòa':'SFNO baseline · loss · giới hạn',
        'Cải thiện nhất quán nhưng có biên nhỏ — đúng kỳ vọng của một residual corrector.':'Validation loss giảm 6,74%; smoke-test cần xem cùng RMSE theo từng kênh.',
        'Chưa có ablation từng thành phần':'Chưa có ablation từng thành phần',
        'Smoke test chỉ 12 samples':'12 mẫu/stage; chưa xác minh khởi tạo ghép cặp',
        'Epoch 0–1':'Epoch 1–2',
        'Epoch 2+':'Epoch 3–4',
        'epoch 3':'epoch 4 (index 3)',
        'epoch 24':'epoch 25 (index 24)',
    }
    for slide in initial:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        for old,new in replacements.items():
                            if old in run.text:run.text=run.text.replace(old,new)

    inserts=[
        (28,'Stage 1 — đường loss huấn luyện','loss1','25 epoch thực tế; best epoch 25 (checkpoint index 24). Train 0,465965 → 0,124181; validation 0,315914 → 0,134443.'),
        (41,'Stage 2 — đường loss fine-tune','loss2','4 epoch thực tế. Epoch 1–2 train corrector; epoch 3–4 mở backbone tail. Best validation: 0,125384.'),
        (41,'Độ ổn định gradient','gradients','Mean/max trước clipping từ history.csv; ngưỡng clip = 1,0. Nonfinite gradient norms = 0 trong các epoch được ghi.'),
        (46,'Sai số theo biến và mức áp suất','pressure','65 pressure channels · 6 surface channels được trình bày riêng. Giá trị dương: Stage 2 giảm RMSE.'),
    ]
    offset=0
    for original_index,heading,key,caption in sorted(inserts,key=lambda item:item[0]):
        slide=standard_slide(prs,logo,0,heading)
        chart_slide(slide,paths[key],caption)
        insert_before(prs,slide,original_index+offset);offset+=1

    apply_wording(prs)
    # Provenance and explicit uncertainty travel with the PPTX speaker notes.
    for index,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.left>Inches(11.7) and (shape.top<Inches(.55) or shape.top>Inches(6.6)):
                if shape.text.strip().isdigit():
                    if shape.text_frame.paragraphs[0].runs:
                        shape.text_frame.paragraphs[0].runs[0].text=f'{index:02d}'
                    else:shape.text=f'{index:02d}'
                    for run in shape.text_frame.paragraphs[0].runs:
                        run.font.size=Pt(10)
                        run.font.color.rgb=rgb(BLUE if shape.top<Inches(.55) else '2F95DE')
        slide.notes_slide.notes_text_frame.text=(
            f'Slide {index}: {title(slide)}\n'
            'Nguồn phương pháp: matching stage source trong submission. '
            'Nguồn kết quả: results/stage_smoke_tests (6 giờ, 12 samples/stage); '
            'train/validation: reports/evidence/stage1_history.csv và stage2_history.csv. '
            'Hai smoke-test có persistence khác nhau do lịch sử đầu vào; chưa chứng minh khởi tạo ghép cặp. '
            'Không dùng các số liệu này để kết luận extreme-event skill hoặc hiệu quả 7 ngày. '
            'Đồ thị phổ ở phần spectral bottleneck là minh họa, không phải spectrum đo từ forecast.'
        )
        # Override template shadows: overlapping theme effects hid alignment errors.
        for shape in slide.shapes:
            for ref in shape._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):
                ref.set('idx','0')
            sppr=shape._element.find('{http://schemas.openxmlformats.org/drawingml/2006/presentation}spPr')
            if sppr is not None:
                prior=sppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst')
                if prior is not None:sppr.remove(prior)
                sppr.append(OxmlElement('a:effectLst'))
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name='Arial'
                        if run.font.size and shape.top>Inches(1) and shape.top<Inches(6.7):
                            current=run.font.size.pt
                            if current<12 and shape.height>=Inches(.25):run.font.size=Pt(12)
        if 46 <= index <= 52:
            text_box(slide,'Nguồn: stage_smoke_tests · lead 6 h · 12 mẫu/stage · chưa xác minh khởi tạo ghép cặp',
                     .85,6.98,11.6,.18,size=9,color=GRAY,align=PP_ALIGN.CENTER)
    return histories,paths

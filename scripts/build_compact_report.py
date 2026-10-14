"""Build a 30-slide report from the reviewed deck and archived rollout evidence.

The original 61-slide deck is retained, so this transformation is repeatable.
No inference is run and no results are synthesized.
"""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
from collections import defaultdict
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.shapes.picture import Picture
from pptx.util import Inches, Pt

from build_full_report_slides import add_brand
from build_report_slides import BLUE, GREEN, NAVY, GRAY, PALE, RED, text_box, rounded_box, add_bullets, rgb

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
ARCHIVE = REPORTS / "archive_61_slides"
EVIDENCE = REPORTS / "evidence" / "rollout_2020"
FIGURES = REPORTS / "figures" / "rollout_2020"
BENCHMARK = ROOT.parent / "NeuralAtmosphereOperator_COMPLETE_EXPORT" / "03_BENCHMARK_AND_REPORTS" / "benchmark_results"
NAME = "NeuralAtmosphereOperator_Full_Report"
NOTES = "NeuralAtmosphereOperator_Presentation_Notes.ipynb"
VARIABLES = ["Z500_gpm", "T850", "T2m", "U10", "V10", "MSLP", "Q850", "U850", "V850"]
LABELS = ["Z500", "T850", "T2M", "U10", "V10", "MSLP", "Q850", "U850", "V850"]
UNITS = ["gpm", "K", "K", "m/s", "m/s", "hPa", "g/kg", "m/s", "m/s"]
LEADS = [6, 12, 24, 72, 120, 168]
COLORS = {"stage1": "#1976D2", "stage2": "#2E8B57", "persistence": "#8794A1", "climatology": "#B57C42"}


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def collect_evidence() -> tuple[list[dict[str, str]], list[dict[str, str]], dict[str, Any]]:
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    allowed = {"stage1", "stage2", "persistence", "climatology"}
    for name in ("milestone_metrics_all_leads.csv", "all_71_channels_all_leads.csv"):
        destination = EVIDENCE / name
        if not destination.exists():
            rows = read_rows(BENCHMARK / "tables" / name)
            selected = [row for row in rows if row["Method"] in allowed]
            with destination.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
                writer.writeheader()
                writer.writerows(selected)
    metadata_path = EVIDENCE / "benchmark_metadata.json"
    if not metadata_path.exists():
        original = json.loads((BENCHMARK / "benchmark_results.json").read_text())
        metadata = {key: original[key] for key in ("timestamp_utc", "test_year", "sample_count", "selected_indices")}
        metadata["spectral_summary"] = {
            lead: {method: data for method, data in values.items() if method in allowed or method == "era5_truth"}
            for lead, values in original["spectral_summary"].items()
        }
        metadata["metric_notes"] = {
            "area_weighting": "cos(latitude), normalized to mean 1 in archived evaluator",
            "acc": "mean spatial anomaly correlation, training climatology",
            "skill": "1 - MSE_model / MSE_baseline; display as percent",
            "confidence_intervals": "95% bootstrap RMSE intervals from archived CSV; not paired improvement intervals",
            "spectral_scope": "Z500 only; spherical degrees 0 through 179",
            "provenance_limit": "selected dataset indices available; exact initialization timestamps and checkpoint hashes not in this benchmark JSON",
            "source": "NeuralAtmosphereOperator_COMPLETE_EXPORT/03_BENCHMARK_AND_REPORTS/benchmark_results",
        }
        metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    metadata = json.loads(metadata_path.read_text())
    rows = read_rows(EVIDENCE / "milestone_metrics_all_leads.csv")
    all_rows = read_rows(EVIDENCE / "all_71_channels_all_leads.csv")
    assert metadata["sample_count"] == len(metadata["selected_indices"]) == 60
    assert {int(row["Lead_Hours"]) for row in rows} == set(LEADS)
    for method in ("stage1", "stage2"):
        assert len([row for row in rows if row["Method"] == method]) == 6
        assert len([row for row in all_rows if row["Method"] == method and row["Lead_Hours"] == "168"]) == 71
    manifest = []
    for path in sorted(EVIDENCE.glob("*")):
        if path.is_file() and path.name != "MANIFEST.sha256":
            manifest.append(f"{hashlib.sha256(path.read_bytes()).hexdigest()}  {path.name}")
    (EVIDENCE / "MANIFEST.sha256").write_text("\n".join(manifest) + "\n", encoding="utf-8")
    return rows, all_rows, metadata


def axis_style(ax: Any) -> None:
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color("#A9B7C6")
    ax.tick_params(labelsize=11, colors="#455564")
    ax.grid(color="#E5EBF1", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.set_xticks([0, 1, 3, 5, 7])


def save_plot(fig: Any, name: str) -> Path:
    FIGURES.mkdir(parents=True, exist_ok=True)
    path = FIGURES / f"{name}.png"
    fig.savefig(path, dpi=240, bbox_inches="tight", facecolor="white")
    fig.savefig(path.with_suffix(".pdf"), bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def plots(rows: list[dict[str, str]], all_rows: list[dict[str, str]], metadata: dict[str, Any]) -> dict[str, Path]:
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 12, "mathtext.fontset": "cm"})
    by = {(row["Method"], int(row["Lead_Hours"])): row for row in rows}
    out: dict[str, Path] = {}
    x = np.array(LEADS) / 24
    for metric in ("RMSE", "ACC", "Bias"):
        fig, axes = plt.subplots(3, 3, figsize=(12.2, 5.65))
        fig.subplots_adjust(left=.06, right=.99, top=.86, bottom=.11, wspace=.28, hspace=.48)
        for ax, variable, label, unit in zip(axes.flat, VARIABLES, LABELS, UNITS):
            methods = ("stage1", "stage2", "persistence") if metric == "ACC" else ("stage1", "stage2")
            for method in methods:
                y = np.array([float(by[method, lead][f"{variable}_{metric}"]) for lead in LEADS])
                ax.plot(x, y, color=COLORS[method], lw=2, marker="o", ms=3.2,
                        ls="--" if method == "persistence" else "-", label=method.replace("stage", "Stage ").capitalize() if method == "persistence" else method.replace("stage", "Stage "))
                if metric == "RMSE":
                    lower = [float(by[method, lead][f"{variable}_RMSE_CI_lo"]) for lead in LEADS]
                    upper = [float(by[method, lead][f"{variable}_RMSE_CI_hi"]) for lead in LEADS]
                    ax.fill_between(x, lower, upper, color=COLORS[method], alpha=.12)
            ax.set_title(label if metric == "ACC" else f"{label} · {unit}", fontsize=13, fontweight="bold", loc="left", color="#20384F", pad=6)
            axis_style(ax)
            if metric == "RMSE":
                ax.set_ylim(bottom=0)
            elif metric == "ACC":
                ax.set_ylim(0, 1.04)
            else:
                ax.axhline(0, color="#8794A1", lw=1, ls=":")
        handles, names = axes[0, 0].get_legend_handles_labels()
        fig.legend(handles, names, loc="upper center", bbox_to_anchor=(.5, 1.01), ncol=3, frameon=False, fontsize=12)
        fig.supxlabel("Thời hạn dự báo (ngày)", fontsize=12, y=.01)
        out[metric.lower()] = save_plot(fig, f"rollout_{metric.lower()}")

    fig, axes = plt.subplots(1, 3, figsize=(12, 4.8))
    fig.subplots_adjust(left=.06, right=.99, top=.8, bottom=.2, wspace=.23)
    for ax, variable, label in zip(axes, ["Z500_gpm", "T2m", "U10"], ["Z500", "T2M", "U10"]):
        lowest_skill = 0.0
        for method in ("stage1", "stage2"):
            for baseline, linestyle in (("persistence", "-"), ("climatology", "--")):
                model_rmse = np.array([float(by[method, lead][f"{variable}_RMSE"]) for lead in LEADS])
                baseline_rmse = np.array([float(by[baseline, lead][f"{variable}_RMSE"]) for lead in LEADS])
                values = 100 * (1 - (model_rmse / baseline_rmse) ** 2)
                lowest_skill = min(lowest_skill, float(values.min()))
                ax.plot(x, values, color=COLORS[method], ls=linestyle, lw=2.2, marker="o", ms=4,
                        label=f"{method.replace('stage', 'Stage ')} / {baseline}")
        axis_style(ax)
        ax.set_ylim(min(-5, lowest_skill - 5), 103)
        ax.axhline(0, color="#8794A1", lw=1)
        ax.set_title(label, loc="left", color="#20384F", fontweight="bold")
    axes[0].set_ylabel("MSE skill (%)")
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", ncol=2, frameon=False, fontsize=11)
    fig.supxlabel("Thời hạn dự báo (ngày)", fontsize=12, y=.08)
    out["skill"] = save_plot(fig, "rollout_baseline_skill")

    fig, axes = plt.subplots(1, 2, figsize=(12, 4.65))
    fig.subplots_adjust(left=.07, right=.99, top=.85, bottom=.17, wspace=.2)
    for ax, key, label in zip(axes, ["06h", "168h_7d"], ["Sau 6 giờ", "Sau 7 ngày"]):
        for method in ("stage1", "stage2"):
            values = np.array(metadata["spectral_summary"][key][method]["R_l"])
            ax.plot(np.arange(1, len(values)), values[1:], lw=2, color=COLORS[method], label=method.replace("stage", "Stage "))
        ax.axhline(1, color="#8794A1", ls="--", lw=1.3, label="ERA5: R = 1")
        ax.set_yscale("log")
        axis_style(ax)
        ax.set_xticks([0, 30, 60, 90, 120, 150, 179])
        ax.set_title(label, loc="left", color="#20384F", fontweight="bold")
        ax.set_xlabel(r"Bậc điều hòa cầu $\ell$")
        ax.legend(frameon=False, fontsize=10)
    axes[0].set_ylabel(r"$R_\ell = P_{model}(\ell) / P_{ERA5}(\ell)$")
    out["spectral"] = save_plot(fig, "rollout_z500_spectral_ratio")

    all_by = {(row["Method"], int(row["Lead_Hours"]), row["Channel_Index"]): row for row in all_rows}
    leads = [6, 24, 72, 120, 168]
    counts, medians = [], []
    groups: dict[str, list[float]] = defaultdict(list)
    for lead in leads:
        changes = []
        for index in range(71):
            a, b = all_by["stage1", lead, str(index)], all_by["stage2", lead, str(index)]
            baseline_rmse = float(a["Physical_RMSE"])
            # The archived CSV rounds Q50 at 6h to zero for both stages.
            # Keep it in absolute channel counts, but never invent a percentage.
            if baseline_rmse == 0:
                continue
            improvement = 100 * (1 - float(b["Physical_RMSE"]) / baseline_rmse)
            changes.append(improvement)
            if lead == 168:
                groups[a["Variable_Group"]].append(improvement)
        counts.append(sum(float(all_by["stage2", lead, str(index)]["Physical_RMSE"]) < float(all_by["stage1", lead, str(index)]["Physical_RMSE"]) for index in range(71)))
        medians.append(float(np.median(changes)))
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.6))
    fig.subplots_adjust(left=.08, right=.98, top=.83, bottom=.2, wspace=.27)
    axes[0].bar(range(5), counts, color=COLORS["stage2"], label="RMSE giảm")
    axes[0].bar(range(5), np.array([71] * 5) - counts, bottom=counts, color="#D94C4C", label="RMSE không giảm")
    axes[0].set_xticks(range(5), ["6 giờ", "1 ngày", "3 ngày", "5 ngày", "7 ngày"])
    axes[0].set_ylim(0, 83)
    axes[0].set_ylabel("Số kênh / 71")
    for i, count in enumerate(counts):
        axes[0].text(i, 73, str(count), ha="center", fontsize=12, fontweight="bold", color="#19663D")
    axes[0].legend(frameon=False, fontsize=10, loc="lower left")
    axes[0].set_title("Kênh có RMSE thấp hơn ở Stage 2", loc="left", color="#20384F", fontsize=13, fontweight="bold")
    axes[1].bar(range(5), medians, color=COLORS["stage2"])
    axes[1].set_xticks(range(5), ["6 giờ", "1 ngày", "3 ngày", "5 ngày", "7 ngày"])
    axes[1].set_ylabel("Mức giảm RMSE trung vị (%)")
    axes[1].set_ylim(0, max(medians) * 1.3)
    for i, value in enumerate(medians):
        axes[1].text(i, value + .08, f"{value:.2f}%", ha="center", fontsize=11, fontweight="bold")
    axes[1].set_title("Trung vị mức giảm RMSE", loc="left", color="#20384F", fontsize=13, fontweight="bold")
    for ax in axes:
        ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="y", color="#E5EBF1")
        ax.set_axisbelow(True)
        ax.tick_params(labelsize=11)
    out["all71"] = save_plot(fig, "rollout_all71")
    return out


def set_text(shape: Any, value: str) -> None:
    frame = shape.text_frame
    first = frame.paragraphs[0]
    if first.runs:
        first.runs[0].text = value
        for run in first.runs[1:]:
            run.text = ""
        for paragraph in frame.paragraphs[1:]:
            paragraph.text = ""
    else:
        shape.text = value


def change(slide: Any, old: str, new: str) -> None:
    for shape in slide.shapes:
        if shape.has_text_frame and old in shape.text:
            set_text(shape, shape.text.replace(old, new))


def picture(slide: Any, path: Path, y: float = 1.14, height: float = 5.35) -> None:
    from PIL import Image
    with Image.open(path) as source:
        width_px, height_px = source.size
    ratio = width_px / height_px
    width = min(12.05, height * ratio)
    height = width / ratio
    slide.shapes.add_picture(str(path), Inches((13.333 - width) / 2), Inches(y), width=Inches(width), height=Inches(height))


def caption(slide: Any, text: str) -> None:
    text_box(slide, text, .65, 6.57, 12.05, .43, size=10.2, color=GRAY, align=PP_ALIGN.CENTER)


def table(slide: Any, headers: list[str], rows: list[list[str]]) -> None:
    widths = [1.12, 1.0, 2.15, 2.15, 2.7, 2.5]
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(.8), Inches(1.55), Inches(sum(widths)), Inches(4.75))
    for column, width in zip(shape.table.columns, widths):
        column.width = Inches(width)
    for i, values in enumerate([headers, *rows]):
        for j, value in enumerate(values):
            cell = shape.table.cell(i, j)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY if i == 0 else ("EEF5FB" if i % 2 else "FFFFFF"))
            cell.margin_top = Inches(.09)
            cell.margin_left = cell.margin_right = Inches(.07)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(13 if i == 0 else 15)
                    run.font.bold = i == 0 or j == 4
                    run.font.color.rgb = rgb("FFFFFF" if i == 0 else GREEN if j == 4 else NAVY)


def build() -> None:
    ARCHIVE.mkdir(parents=True, exist_ok=True)
    for filename in (f"{NAME}.pptx", f"{NAME}.pdf", NOTES):
        source, target = REPORTS / filename, ARCHIVE / filename
        if source.exists() and not target.exists():
            shutil.copyfile(source, target)
    prs = Presentation(str(ARCHIVE / f"{NAME}.pptx"))
    assert len(prs.slides) == 61, "Expected the preserved 61-slide source deck"
    original = {number: slide for number, slide in enumerate(prs.slides, 1)}
    identifiers = {slide.part.partname: node for slide, node in zip(prs.slides, list(prs.slides._sldIdLst))}
    notes = json.loads((ARCHIVE / NOTES).read_text())
    logo_shape = next(shape for shape in original[25].shapes if shape.shape_type == 13)
    assert isinstance(logo_shape, Picture)
    logo = REPORTS / "figures" / "hcmute_logo.png"
    logo.write_bytes(logo_shape.image.blob)
    rows, all_rows, metadata = collect_evidence()
    figures = plots(rows, all_rows, metadata)
    by = {(row["Method"], int(row["Lead_Hours"])): row for row in rows}
    ordered: list[Any] = []
    speeches: list[tuple[str, int, str, str]] = []
    sources: list[str] = []

    def keep(number: int, script: str | None = None, seconds: int | None = None, cue: str = "") -> Any:
        slide = original[number]
        ordered.append(slide)
        cell = notes["cells"][number - 1]
        title = cell["source"][0].split(" — ", 1)[1].strip()
        speeches.append((title, seconds or cell["metadata"]["target_seconds"], script or cell["source"][4].strip(), cue))
        sources.append(f"Ancien slide {number}")
        return slide

    def fresh(title: str, script: str, seconds: int = 20, cue: str = "") -> Any:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_brand(slide, logo, title, len(ordered) + 1)
        ordered.append(slide)
        speeches.append((title, seconds, script, cue))
        sources.append("Nouveau / synthèse")
        return slide

    keep(1)
    slide = keep(4, "Đầu vào là state hiện tại; Stage 2 sử dụng thêm state trước đó sáu giờ. Đầu ra gồm 71 trường tại thời điểm sau sáu giờ. Khi rollout, model dùng dự báo của chính nó cho bước tiếp theo. Báo cáo có cả đường loss khi train và benchmark đến bảy ngày.")
    change(slide, "Giảm sai số dự báo 6 giờ; theo dõi cấu trúc không gian và độ ổn định khi dự báo lặp.", "Huấn luyện bước 6 giờ; đánh giá dự báo lặp đến 7 ngày trên tập test.")
    keep(9)
    slide = keep(10, "71 kênh gồm sáu biến bề mặt và năm nhóm biến trên 13 mức áp suất. Năm nhóm là hai thành phần gió, nhiệt độ, độ ẩm riêng và geopotential. Geopotential là thế năng trọng trường trên mỗi kilogram. Các biến có đơn vị khác nhau, nên cần chuẩn hóa trước khi train.")
    change(slide, "Địa thế vị", "Geopotential")
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text == "50 → 1000 hPa" and shape.left < Inches(2) and shape.top > Inches(4):
            set_text(shape, "Thế năng trọng trường trên mỗi kg\n50–1000 hPa · m²/s²")
            shape.top = Inches(4.65)
            shape.height = Inches(.85)
    slide = fresh("Chia dữ liệu và hai phép kiểm thử", "Train dùng dữ liệu 1983 đến 2018; validation dùng năm 2019 để chọn checkpoint. Benchmark gốc ghi 60 khởi tạo trong tập test năm 2020, rollout đến bảy ngày. Một phép kiểm thử riêng dùng 12 khởi tạo trong tháng bảy năm 2020 ở sáu giờ. Hai phép đánh giá này không được gộp số liệu.")
    for title, body, x in [("TRAIN", "1983–2018\nTrain và tính thống kê", .65), ("VALIDATION", "2019\nChọn checkpoint", 4.85), ("TEST", "2020\nBenchmark lưu trữ", 9.05)]:
        rounded_box(slide, x, 1.55, 3.6, 1.6, fill=PALE, line=BLUE)
        text_box(slide, title, x+.2, 1.77, 3.2, .35, size=19, bold=True, color=BLUE)
        text_box(slide, body, x+.2, 2.3, 3.2, .7, size=16)
    add_bullets(slide, ["Benchmark gốc: 60 khởi tạo · rollout 6 giờ → 7 ngày · Stage 1 và Stage 2", "Smoke test riêng: tháng 7/2020 · 12 khởi tạo/stage · thời hạn 6 giờ", "Thống kê chuẩn hóa và climatology lấy từ tập train, không từ test"], .9, 3.8, 11.6, 1.9, size=17, spacing=16)
    keep(13)
    slide = fresh("Đọc các chỉ số đánh giá", "RMSE đo độ lớn sai số, ACC đo mức tương đồng của dị thường so với climatology. Bias cho biết lệch trung bình theo hướng nào. MSE skill so với baseline dương là tốt hơn; còn tỷ số năng lượng phổ gần một nghĩa là năng lượng gần ERA5 hơn, nhưng chưa chứng minh dự báo đúng vị trí.")
    descriptions = [("RMSE ↓", "Sai số theo đơn vị vật lý; không so trực tiếp giữa các biến"), ("ACC ↑", "Tương quan không gian của dị thường; gần 1 là tốt hơn"), ("Bias → 0", "Lệch trung bình có dấu; dương và âm có thể triệt tiêu"), ("MSE skill ↑", "1 − MSE model / MSE baseline; không phải % giảm RMSE"), ("Spectral ratio → 1", "Năng lượng phổ model / ERA5 tại từng bậc điều hòa cầu")]
    for i, (title, body) in enumerate(descriptions):
        y = 1.24 + i * .98
        rounded_box(slide, .75, y, 11.85, .8, fill=PALE, line="D3DEE8")
        text_box(slide, title, 1, y+.2, 2.6, .35, size=18, color=BLUE, bold=True)
        text_box(slide, body, 3.6, y+.2, 8.7, .4, size=16)
    caption(slide, "Variance ratio không gian khác spectral ratio; benchmark rollout không cung cấp variance ratio.")
    keep(17)
    keep(19)
    slide = keep(21)
    change(slide, "Chia sai số cho độ lệch chuẩn của tendency", "Sai số state chuẩn hóa chia cho tendency scale s")
    slide = keep(25)
    text_box(slide, "scale_factor=2 · lưới phổ nội bộ 181 × 360", .9, 6.48, 11.5, .3, size=14, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    slide = keep(27)
    change(slide, "Lịch learning rate", "Learning rate schedule")
    keep(29, cue="Chỉ đường xanh dương (train), xanh lá (validation) và điểm epoch 25.")
    slide = keep(36)
    caption(slide, "Loss tính trên dự báo sau khi cộng correction; corrector không dùng loss riêng.")
    slide = keep(37, "Bảy kênh gồm địa hình, mặt nạ đất–biển, cosine góc thiên đỉnh Mặt Trời, sin–cos giờ địa phương và pha trong năm. Chúng được concatenate vào đầu vào corrector, trước encoder. Backbone vẫn chỉ nhận state hiện tại. Cosine góc thiên đỉnh là thông tin hình học, không phải bức xạ đo được.")
    change(slide, "Forcing chỉ đi vào corrector", "7 kênh forcing được concatenate trước encoder của corrector")
    slide = keep(39, "Backbone nhận 71 kênh hiện tại. Corrector ghép state, tendency gần nhất và tendency từ backbone, mỗi nhóm 71 kênh, thêm bảy forcing và hai latitude features: tổng cộng 222. Encoder chuyển thành 84 kênh, qua bốn khối SFNO trên lưới gốc, rồi trả về correction 71 kênh.", cue="Chỉ nhánh backbone, rồi phép concatenate và corrector.")
    text_box(slide, "Corrector: scale_factor=1 · 361 × 720 · decoder khởi tạo bằng 0", .8, 6.68, 11.8, .25, size=12, color=GRAY, align=PP_ALIGN.CENTER)
    keep(41)
    keep(43, cue="Chỉ hai pha train; phân biệt validation loss với RMSE trên test.")

    slide = fresh("Benchmark rollout: 28 bước đến 7 ngày", "Benchmark lưu trữ đánh giá 60 khởi tạo trong tập test năm 2020. Model tiến từng bước sáu giờ, không nhận state thực tế giữa chừng, và được chấm tại sáu thời hạn đến 168 giờ. Stage 1, Stage 2, persistence và climatology dùng chung tập khởi tạo trong runner của benchmark.")
    timeline = [("6 giờ", "1 bước"), ("12 giờ", "2 bước"), ("1 ngày", "4 bước"), ("3 ngày", "12 bước"), ("5 ngày", "20 bước"), ("7 ngày", "28 bước")]
    for i, (a, b) in enumerate(timeline):
        x0 = .6 + i * 2.12
        rounded_box(slide, x0, 1.55, 1.83, 1.22, fill=PALE, line=BLUE)
        text_box(slide, a, x0+.12, 1.76, 1.6, .35, size=20, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
        text_box(slide, b, x0+.12, 2.25, 1.6, .25, size=14, align=PP_ALIGN.CENTER)
    add_bullets(slide, ["Autoregressive rollout: dự báo bước trước trở thành đầu vào bước sau", "Model và baselines được chấm trên cùng khởi tạo trong benchmark gốc", "Trọng số cos(vĩ độ); ACC trừ climatology từ tập train", "Dải tô trên RMSE: khoảng tin cậy bootstrap 95% đã lưu trong CSV"], .9, 3.4, 11.55, 2.4, size=18, spacing=17)
    caption(slide, "Nguồn: benchmark gốc ngày 15/09/2026 · 60 khởi tạo · không phải smoke test 12 mẫu.")

    slide = fresh("RMSE theo thời hạn dự báo", "Mỗi ô là một biến với đơn vị và trục riêng. Đường xanh dương là Stage 1, xanh lá là Stage 2; dải tô là khoảng tin cậy RMSE đã lưu. Sai số tăng theo thời hạn. Với Z500, Stage 2 tăng từ 3,39 gpm ở sáu giờ lên 72,67 gpm ở bảy ngày, nhưng vẫn thấp hơn Stage 1. Dải tin cậy của từng model không phải khoảng tin cậy cho chênh lệch giữa hai model.", 35, "Theo dõi Z500 từ 6 giờ đến 7 ngày; chỉ thêm ô T2M.")
    picture(slide, figures["rmse"])
    caption(slide, "Benchmark 60 khởi tạo · test 2020 · 6–168 giờ · dải tô: bootstrap 95% cho RMSE từng model.")
    slide = fresh("ACC theo thời hạn dự báo", "ACC giảm theo thời hạn, nhưng mức giảm khác nhau giữa các biến. Sau bảy ngày, Stage 2 có ACC Z500 khoảng 0,769 và T2M khoảng 0,877, trong khi U10 chỉ còn 0,465. Model giữ cấu trúc dị thường tốt hơn ở một số biến so với gió bề mặt. ACC cao ở nhiệt độ không đủ để kết luận mọi trường khí quyển đều dự báo tốt.", 35, "Chỉ Z500, T2M và U10; đường xám là persistence.")
    picture(slide, figures["acc"])
    caption(slide, "ACC dùng climatology từ tập train · đường xám: persistence · không đặt một ngưỡng chất lượng chung cho mọi biến.")

    day7_rows = []
    for variable, label, unit in zip(VARIABLES, LABELS, UNITS):
        a = float(by["stage1", 168][f"{variable}_RMSE"])
        b = float(by["stage2", 168][f"{variable}_RMSE"])
        acc = float(by["stage2", 168][f"{variable}_ACC"])
        day7_rows.append([label, unit, f"{a:.3f}", f"{b:.3f}", f"{100*(1-b/a):.2f}%", f"{acc:.4f}"])
    slide = fresh("Kết quả tại 7 ngày: Stage 1 và Stage 2", "Tại bảy ngày, RMSE của Z500 giảm từ 75,01 xuống 72,67 gpm, tương ứng 3,12%. T2M giảm từ 2,753 xuống 2,663 kelvin; U10 giảm từ 4,364 xuống 4,232 mét trên giây. Cả chín biến trong bảng có RMSE thấp hơn ở Stage 2. Cột cuối là ACC của Stage 2, không phải phần trăm cải thiện. Các mức giảm này mô tả benchmark lưu trữ, chưa phải kiểm định chênh lệch ghép cặp.", 40, "Chỉ ba hàng Z500, T2M, U10; không đọc cả bảng.")
    table(slide, ["Biến", "Đơn vị", "RMSE Stage 1", "RMSE Stage 2", "RMSE giảm", "ACC Stage 2"], day7_rows)
    caption(slide, "Z500: độ cao quy đổi từ geopotential tại 500 hPa (gpm) · thời hạn 168 giờ · 60 khởi tạo.")

    slide = fresh("MSE skill so với persistence và climatology", "Skill dương nghĩa là MSE thấp hơn baseline; 50% là giảm MSE một nửa, không phải giảm RMSE một nửa. Nét liền so với persistence, nét đứt so với climatology. Với U10 ở ngày bảy, cả hai stage tốt hơn persistence nhưng kém climatology; Stage 2 có skill khoảng âm 4,7% so với climatology. Vì vậy, chất lượng cần được đối chiếu với nhiều baseline, không chỉ persistence.", 30, "Chỉ chú giải nét liền/nét đứt và U10 ở ngày 7.")
    picture(slide, figures["skill"], y=1.35, height=4.9)
    text_box(slide, "U10 tại 7 ngày: Stage 2 tốt hơn persistence, nhưng kém climatology (skill −4,7%).", .9, 6.03, 11.6, .33, size=14, color=RED, bold=True, align=PP_ALIGN.CENTER)
    caption(slide, "Nét liền: so với persistence · nét đứt: so với climatology · dương = MSE thấp hơn baseline.")
    slide = fresh("Bias khi rollout: sai lệch trung bình có dấu", "Bias gần không nghĩa là sai lệch trung bình nhỏ, không phải mọi điểm đều chính xác. Sai số dương và âm có thể triệt tiêu nhau. Các đường trên slide cho biết bias thay đổi theo thời hạn và theo biến. Cần đọc cùng RMSE và ACC để phân biệt sai lệch hệ thống với sai số không gian; không thể dùng riêng bias để kết luận rollout ổn định.", 30, "Chỉ đường 0 và một biến có bias đổi theo thời hạn.")
    picture(slide, figures["bias"])
    caption(slide, "Bias = trung bình có trọng số của dự báo − thực tế · mỗi ô có đơn vị riêng · gần 0 không đồng nghĩa RMSE thấp.")
    slide = fresh("Phổ Z500: năng lượng không được giữ đồng đều", "Đây là phổ Z500 đo trong benchmark, không phải hình minh họa. Tỷ số bằng một nghĩa là năng lượng model bằng ERA5 tại bậc đó. Ở ngày bảy có vùng thiếu năng lượng và vùng dư năng lượng; tại bậc 179, tỷ số khoảng 13,8 ở Stage 1 và 11,8 ở Stage 2. Vì năng lượng tham chiếu có thể nhỏ, cần xem thêm phổ tuyệt đối. Dữ liệu phổ chỉ tới bậc 179, chưa kiểm tra toàn bộ dải tần của lưới gốc.", 40, "Chỉ đường R=1, vùng dưới 1 và phần tăng ở bậc cao; trục dọc là log.")
    picture(slide, figures["spectral"], y=1.4, height=4.9)
    caption(slide, "Z500 · Rℓ = Pmodel / PERA5 · trục log · chỉ ℓ=1…179; không suy rộng sang toàn bộ 71 biến hoặc mọi tần số.")
    slide = fresh("Kiểm tra trên toàn bộ 71 kênh", "Trong benchmark 60 khởi tạo, Stage 2 giảm RMSE ở 68 trên 71 kênh tại sáu giờ và 65 kênh tại bảy ngày. Như vậy vẫn có các kênh không cải thiện. Biểu đồ bên phải là mức giảm trung vị, giúp tránh chỉ nhìn vài biến ưu tiên. Những số này thuộc benchmark gốc; chúng khác bảng smoke test 12 khởi tạo vì hai phép kiểm thử dùng tập mẫu khác nhau.", 30, "Chỉ số kênh tại 6 giờ và 7 ngày, sau đó mức giảm trung vị.")
    picture(slide, figures["all71"], y=1.4, height=4.85)
    caption(slide, "60 khởi tạo · trung vị 6 giờ bỏ Q50 vì RMSE làm tròn bằng 0 (70 kênh); các mốc còn lại: 71 kênh.")

    slide = fresh("Giới hạn và phạm vi kết luận", "Model đã được test đến bảy ngày, dù train một bước. Benchmark có 60 khởi tạo, chưa đánh giá riêng cực trị và chưa có ablation. Metadata lưu chỉ số mẫu nhưng thiếu timestamp khởi tạo và hash checkpoint. Phổ chỉ có Z500 tới bậc 179. Những giới hạn này cần được nêu khi giải thích kết quả.")
    add_bullets(slide, ["Đã rollout đến 7 ngày; train K=1 không có nghĩa là chỉ test K=1", "60 khởi tạo; chưa đánh giá riêng bão, cực trị và tính đại diện theo mùa", "Chưa có ablation study hoặc khoảng tin cậy cho chênh lệch ghép cặp", "Benchmark JSON chưa lưu đầy đủ timestamp khởi tạo và hash checkpoint", "Phổ lưu trữ chỉ có Z500, ℓ≤179; chưa xác nhận chất lượng toàn dải", "Không so trực tiếp với paper dùng năm test, lưới hoặc climatology khác"], .85, 1.4, 11.7, 4.8, size=18, spacing=17)
    caption(slide, "Kết quả lưu trữ là bằng chứng thực nghiệm; không thay thế một lượt đánh giá được kiểm chứng lại đầy đủ.")
    slide = fresh("Thí nghiệm tiếp theo: ưu tiên kiểm chứng", "Em ưu tiên chạy lại benchmark với metadata đầy đủ, rồi bỏ từng thành phần để làm ablation. Sau đó thử multi-step training, spectral loss hoặc tăng trọng số cực trị. Mỗi thay đổi cần đánh giá theo thời hạn và sự kiện, đồng thời theo dõi chi phí. Đây là đề xuất, chưa phải kết quả đã đạt được.")
    for title, body, x in [("1. Đánh giá", "Cùng khởi tạo, hash checkpoint\nPhân tích theo mùa và sự kiện\nSai số + khoảng tin cậy", .65), ("2. Ablation", "Bỏ history / forcing\nCorrector lưới gốc / giảm lưới\nBackbone cố định / mở khóa", 4.85), ("3. Loss và train", "Multi-step training\nSpectral / gradient loss\nSampling và trọng số cực trị", 9.05)]:
        rounded_box(slide, x, 1.7, 3.6, 3.9, fill=PALE, line=BLUE)
        text_box(slide, title, x+.18, 1.99, 3.2, .4, size=21, color=BLUE, bold=True)
        text_box(slide, body, x+.18, 2.87, 3.22, 2.3, size=17)
    caption(slide, "Đổi từng yếu tố; ghi rõ chi phí và mục tiêu đo; không mặc định tần số cao là nhiễu.")
    slide = fresh("Kết luận", "Em đã xây dựng SFNO baseline và spherical residual corrector. Stage 2 giảm validation loss 6,74%. Trong benchmark bảy ngày, RMSE Z500 giảm 3,12% và 65 trên 71 kênh có RMSE thấp hơn Stage 1. Cải thiện không đồng đều; phổ Z500 vẫn có sai lệch năng lượng. Hướng tiếp theo là kiểm chứng benchmark, làm ablation và đánh giá cực trị, không chỉ tiếp tục giảm loss trung bình.", 30)
    cards = [("Validation loss", "−6,74%", "Checkpoint Stage 2 so với Stage 1"), ("Z500 RMSE · 7 ngày", "−3,12%", "75,012 → 72,674 gpm"), ("Kênh RMSE giảm · 7 ngày", "65 / 71", "Benchmark gốc · 60 khởi tạo")]
    for i, (label, value, detail) in enumerate(cards):
        x0 = .65 + i * 4.2
        rounded_box(slide, x0, 1.6, 3.6, 2.45, fill=PALE, line=BLUE)
        text_box(slide, label, x0+.18, 1.89, 3.2, .65, size=17, bold=True, color=BLUE)
        text_box(slide, value, x0+.18, 2.6, 3.2, .6, size=32, bold=True, color=GREEN)
        text_box(slide, detail, x0+.18, 3.4, 3.2, .42, size=12.5, color=GRAY)
    add_bullets(slide, ["Cải thiện không đồng đều; RMSE giảm không chứng minh phổ được phục hồi", "Đã test đến 7 ngày, nhưng chưa đánh giá riêng các sự kiện cực đoan", "Cần ablation study và provenance đầy đủ cho lượt đánh giá tiếp theo"], .95, 4.7, 11.5, 1.7, size=18, spacing=14)
    keep(61)
    assert len(ordered) == len(speeches) == 30
    chosen = {slide.part.partname for slide in ordered}
    for node in list(prs.slides._sldIdLst):
        if prs.part.related_slide(node.rId).part.partname not in chosen:
            prs.part.drop_rel(node.rId)
        prs.slides._sldIdLst.remove(node)
    for number, slide in enumerate(ordered, 1):
        node = identifiers.get(slide.part.partname)
        if node is None:
            # Added slides keep their relationship even after ordering nodes are removed.
            relationship = next(rel for rel in prs.part.rels.values() if rel.reltype.endswith('/slide') and rel.target_part is slide.part)
            from pptx.oxml.xmlchemy import OxmlElement
            node = OxmlElement("p:sldId")
            node.set("id", str(1000 + number))
            node.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", relationship.rId)
        prs.slides._sldIdLst.append(node)
        # Standard page number or cover/Q&A number, preserving its run formatting.
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip().isdigit() and len(shape.text.strip()) <= 2 and (shape.left > Inches(11) or shape.top > Inches(6)):
                set_text(shape, f"{number:02d}")
        title, seconds, script, cue = speeches[number - 1]
        slide.notes_slide.notes_text_frame.text = f"Thời lượng dự kiến: {seconds} giây\n\n{script}" + (f"\n\nThao tác: {cue}" if cue else "")
    prs.save(str(REPORTS / f"{NAME}.pptx"))
    cells = []
    for number, (title, seconds, script, cue) in enumerate(speeches, 1):
        source = [f"## Slide {number:02d} — {title}\n", "\n", f"**Thời lượng dự kiến: {seconds} giây**\n", "\n", script + "\n"]
        if cue:
            source.extend(["\n", f"*Thao tác khi nói (không đọc): {cue}*\n"])
        cells.append({"cell_type": "markdown", "id": f"slide-{number:02d}", "metadata": {"slide_number": number, "target_seconds": seconds}, "source": source})
    total = sum(item[1] for item in speeches)
    notebook = {"cells": cells, "metadata": {"language_info": {"name": "markdown"}, "presentation": {"slide_count": 30, "total_target_seconds": total, "source": f"reports/{NAME}.pptx", "benchmark": "60 initializations, test 2020, leads 6 to 168 hours", "timing_note": "Thời lượng mục tiêu luyện nói, chưa tính chuyển slide; chỉ đọc đoạn lời trình bày."}}, "nbformat": 4, "nbformat_minor": 5}
    (REPORTS / NOTES).write_text(json.dumps(notebook, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    mapping = ["# Bản báo cáo rút gọn — 30 slide", "", "Bản cũ 61 slide được giữ tại `archive_61_slides/`. Bản mới giữ kiến trúc, loss, training curves và bổ sung benchmark rollout đến 7 ngày. Notebook và speaker notes khớp thứ tự slide.", "", f"Thời lượng mục tiêu: {total//60} phút {total%60:02d} giây, chưa tính chuyển slide.", "", "## Thay đổi", "", "- Bỏ slide chuyển phần, mục lục dài, bảng kích thước tensor riêng, cấu hình lặp và danh sách hướng nghiên cứu dàn trải.", "- Gộp thông tin lưới Stage 1 vào kiến trúc; corrector E84/L4 và lưới gốc vào sơ đồ Stage 2.", "- Thay phần kết quả tập trung smoke test bằng benchmark gốc 60 khởi tạo, 6–168 giờ. Không trộn hai phép kiểm thử.", "- Sửa nhận định sai rằng chưa test rollout dài; train K=1 vẫn có thể test autoregressive 28 bước.", "- Phổ thực nghiệm Z500 có vùng dư/thiếu năng lượng; không dùng nhận định cũ rằng không suy hao hay nổ phổ.", "- Không dùng các giá trị training loss từng bị gắn nhãn validation trong báo cáo benchmark cũ; validation lấy từ history đã kiểm tra.", "", "## Thứ tự mới", "", "| Slide | Nội dung | Nguồn | Giây |", "|---:|---|---|---:|"]
    for index, ((title, seconds, _, _), origin) in enumerate(zip(speeches, sources), 1):
        mapping.append(f"| {index} | {title} | {origin.replace('Ancien slide', 'Slide cũ').replace('Nouveau / synthèse', 'Mới / gộp')} | {seconds} |")
    mapping.extend(["", "## Phạm vi bằng chứng", "", "Các CSV tại `evidence/rollout_2020/` được lọc từ bản export, chỉ giữ Stage 1, Stage 2 và baselines. Dữ liệu gốc không bị sửa. Biểu đồ tính từ CSV và spectral arrays lưu trữ, không chạy inference mới. Tỷ số phổ chỉ có Z500 và bậc 0–179. MSE skill được tính lại từ RMSE model và baseline, hiển thị phần trăm. CI là CI của từng model, chưa phải CI cho chênh lệch ghép cặp. Metadata gốc lưu test year và chỉ số mẫu, chưa có đầy đủ timestamp/hash checkpoint. Evaluator cũ dùng cos(vĩ độ), không phải exact latitude-cell-boundary weights như evaluator mới.", ""])
    (REPORTS / "COMPACT_REPORT_REVIEW.md").write_text("\n".join(mapping), encoding="utf-8")
    print(f"Built 30 slides and 30 Markdown cells; target duration {total//60}:{total%60:02d}.")


if __name__ == "__main__":
    build()
